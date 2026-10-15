# converter_excel_para_csv.py
# Converte a(s) Matriz(es) de Recolha de Informação (.xlsx) em ficheiros .csv,
# um por província, prontos a ser lidos pelo importar_dados.py.
#
# Esta ferramenta só é precisa quando chegar uma matriz NOVA em Excel (ex.:
# de uma província que ainda não temos, ou uma atualização de uma província
# já existente). O site em si (e o deploy no Render) já não precisa de abrir
# nenhum ficheiro .xlsx — só lê os .csv gerados aqui.
#
# Uso:
#   python converter_excel_para_csv.py                      -> converte tudo o que estiver em origem_excel/
#   python converter_excel_para_csv.py caminho/ficheiro.xlsx -> converte um único ficheiro
#   python converter_excel_para_csv.py caminho/pasta         -> converte todos os .xlsx dessa pasta
#
# O resultado fica em data/<PROVINCIA>.csv — se já existir um CSV dessa
# província, é substituído (para não duplicar registos).

import sys
import os
import csv
import glob
import openpyxl

from tagging import normalizar, detetar_provincia, detetar_distritos, mapear_distrito_por_nome_aproximado

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PASTA_ORIGEM_PADRAO = os.path.join(BASE_DIR, 'origem_excel')
PASTA_DESTINO_CSV = os.path.join(BASE_DIR, 'data')

LINHAS_PROCURAR_CABECALHO = 30
COL_NUMERO = 0

CAMPOS_CSV = [
    'numero_original', 'nome', 'natureza', 'origem', 'provincia',
    'area_intervencao', 'programas', 'projetos_em_curso', 'fonte_financiamento',
    'distritos_texto', 'distritos_extra', 'perfil_grupo_alvo', 'pessoa_contacto', 'telefone', 'email',
    'coordenada_x', 'coordenada_y',
]


def limpar(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    valor = str(valor).strip()
    return valor


def localizar_linha_cabecalho(sheet):
    for r in range(1, min(LINHAS_PROCURAR_CABECALHO, sheet.max_row) + 1):
        for c in range(1, 4):
            v = sheet.cell(row=r, column=c).value
            if v and 'designacao' in normalizar(v):
                return r
    return None


def localizar_texto_provincia_de(sheet, ate_linha):
    for r in range(1, ate_linha):
        for c in range(1, 4):
            v = sheet.cell(row=r, column=c).value
            if v and 'provincia de:' in normalizar(v):
                return v
    return None


def construir_mapa_colunas(sheet, linha_principal, linha_sub):
    max_col = sheet.max_column
    principal = [sheet.cell(row=linha_principal, column=c).value for c in range(1, max_col + 1)]
    sub = [sheet.cell(row=linha_sub, column=c).value for c in range(1, max_col + 1)]

    principal_ff = []
    ultimo = None
    for v in principal:
        if v:
            ultimo = v
        principal_ff.append(ultimo)

    mapa = {}
    for idx in range(max_col):
        if idx == COL_NUMERO:
            continue
        texto_sub = normalizar(sub[idx]) if sub[idx] else ''
        texto_principal = normalizar(principal_ff[idx]) if principal_ff[idx] else ''

        if texto_sub:
            if texto_sub == 'ong':
                mapa['natureza_ong'] = idx
            elif 'osc' in texto_sub:
                mapa['natureza_osc'] = idx
            elif 'fundac' in texto_sub:
                mapa['natureza_fundacao'] = idx
            elif texto_sub == 'nacional':
                mapa['origem_nacional'] = idx
            elif texto_sub in ('internacional', 'estrangeira'):
                mapa['origem_internacional'] = idx
            elif texto_sub == 'x':
                mapa['coordenada_x'] = idx
            elif texto_sub == 'y':
                mapa['coordenada_y'] = idx
            elif texto_sub == 'telefone':
                mapa['telefone'] = idx
            elif 'mail' in texto_sub:
                mapa['email'] = idx
            elif texto_sub == 'provincia':
                mapa['provincia_coluna'] = idx
            elif 'distrito' in texto_sub or 'municipio' in texto_sub:
                mapa['distritos_texto'] = idx
            elif texto_sub.startswith('programa'):
                mapa['programas'] = idx
            elif texto_sub.startswith('projec') or texto_sub.startswith('projet'):
                mapa['projetos_em_curso'] = idx
        elif texto_principal:
            if 'designacao' in texto_principal:
                mapa['nome'] = idx
            elif 'area' in texto_principal and ('intervenc' in texto_principal or 'tematic' in texto_principal):
                mapa['area_intervencao'] = idx
            elif 'distrito' in texto_principal:
                mapa['distritos_texto'] = idx
            elif 'projec' in texto_principal or 'projet' in texto_principal:
                mapa['projetos_em_curso'] = idx
            elif 'financiamento' in texto_principal:
                mapa['fonte_financiamento'] = idx
            elif 'perfil' in texto_principal:
                mapa['perfil_grupo_alvo'] = idx
            elif 'ponto focal' in texto_principal or ('pessoa' in texto_principal and 'contacto' in texto_principal):
                mapa['pessoa_contacto'] = idx

    return mapa


def valor(row, mapa, campo):
    idx = mapa.get(campo)
    if idx is None or idx >= len(row):
        return ''
    return limpar(row[idx])


def montar_natureza(row, mapa):
    partes = []
    if valor(row, mapa, 'natureza_ong'):
        partes.append('ONG')
    if valor(row, mapa, 'natureza_osc'):
        partes.append('OSC')
    if valor(row, mapa, 'natureza_fundacao'):
        partes.append('Fundação')
    return ', '.join(partes)


def montar_origem(row, mapa):
    nacional = valor(row, mapa, 'origem_nacional')
    internacional = valor(row, mapa, 'origem_internacional')
    if nacional and internacional:
        return 'Nacional e Estrangeira'
    if nacional:
        return 'Nacional'
    if internacional:
        return 'Estrangeira'
    return ''


def converter_folha(sheet, nome_ficheiro, provincia_do_ficheiro=None):
    """Devolve uma lista de dicts (um por registo/linha de ONG) desta folha."""
    linha_principal = localizar_linha_cabecalho(sheet)
    if not linha_principal:
        return []

    linha_sub = linha_principal + 1
    primeira_linha_dados = linha_principal + 2
    mapa = construir_mapa_colunas(sheet, linha_principal, linha_sub)

    texto_provincia_celula = localizar_texto_provincia_de(sheet, linha_principal)
    provincia_da_celula = detetar_provincia(texto_provincia_celula) if texto_provincia_celula else None
    provincia_da_folha = provincia_do_ficheiro or provincia_da_celula

    if provincia_da_celula and provincia_do_ficheiro and provincia_da_celula != provincia_do_ficheiro:
        print(f'  AVISO: "{nome_ficheiro}" (folha "{sheet.title}") - o nome do ficheiro indica '
              f'"{provincia_do_ficheiro}" mas a célula interna diz "{provincia_da_celula}". '
              f'A usar "{provincia_do_ficheiro}" (nome do ficheiro).')

    distrito_da_folha = mapear_distrito_por_nome_aproximado(sheet.title, provincia_da_folha) \
        if provincia_da_folha else None

    registos = []
    for row in sheet.iter_rows(min_row=primeira_linha_dados, values_only=True):
        nome = valor(row, mapa, 'nome') if mapa.get('nome') is not None else ''
        if not nome or normalizar(nome) == 'legenda':
            continue
        if len(row) <= COL_NUMERO or not limpar(row[COL_NUMERO]):
            continue

        provincia_coluna_texto = valor(row, mapa, 'provincia_coluna')
        provincia_da_linha = (detetar_provincia(provincia_coluna_texto) if provincia_coluna_texto else None) \
            or provincia_da_folha
        if not provincia_da_linha:
            continue

        distritos_texto_linha = valor(row, mapa, 'distritos_texto') or (distrito_da_folha or '')
        distritos_extra = ''
        if distrito_da_folha:
            detectados = set(detetar_distritos(distritos_texto_linha, provincia_da_linha))
            if distrito_da_folha not in detectados:
                distritos_extra = distrito_da_folha

        registos.append({
            'numero_original': limpar(row[COL_NUMERO]),
            'nome': nome,
            'natureza': montar_natureza(row, mapa),
            'origem': montar_origem(row, mapa),
            'provincia': provincia_da_linha,
            'area_intervencao': valor(row, mapa, 'area_intervencao'),
            'programas': valor(row, mapa, 'programas'),
            'projetos_em_curso': valor(row, mapa, 'projetos_em_curso'),
            'fonte_financiamento': valor(row, mapa, 'fonte_financiamento'),
            'distritos_texto': distritos_texto_linha,
            'distritos_extra': distritos_extra,
            'perfil_grupo_alvo': valor(row, mapa, 'perfil_grupo_alvo'),
            'pessoa_contacto': valor(row, mapa, 'pessoa_contacto'),
            'telefone': valor(row, mapa, 'telefone'),
            'email': valor(row, mapa, 'email'),
            'coordenada_x': valor(row, mapa, 'coordenada_x'),
            'coordenada_y': valor(row, mapa, 'coordenada_y'),
        })
    return registos


def converter_ficheiro(caminho_xlsx):
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    nome_ficheiro = os.path.basename(caminho_xlsx)
    provincia_do_ficheiro = detetar_provincia(os.path.splitext(nome_ficheiro)[0])

    registos_por_provincia = {}
    for sheet in wb.worksheets:
        registos = converter_folha(sheet, nome_ficheiro, provincia_do_ficheiro)
        for registo in registos:
            registos_por_provincia.setdefault(registo['provincia'], []).append(registo)

    return registos_por_provincia, nome_ficheiro


def gravar_csv(provincia, registos):
    os.makedirs(PASTA_DESTINO_CSV, exist_ok=True)
    nome_csv = normalizar(provincia).replace(' ', '_').upper() + '.csv'
    caminho = os.path.join(PASTA_DESTINO_CSV, nome_csv)
    with open(caminho, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=CAMPOS_CSV)
        writer.writeheader()
        writer.writerows(registos)
    return caminho


def listar_ficheiros(caminho):
    if os.path.isdir(caminho):
        return sorted(glob.glob(os.path.join(caminho, '*.xlsx')))
    if os.path.isfile(caminho):
        return [caminho]
    return []


def converter(caminho):
    ficheiros = listar_ficheiros(caminho)
    if not ficheiros:
        print(f'Nenhum ficheiro .xlsx encontrado em: {caminho}')
        sys.exit(1)

    todos_registos = {}
    for ficheiro in ficheiros:
        registos_por_provincia, nome_ficheiro = converter_ficheiro(ficheiro)
        for provincia, registos in registos_por_provincia.items():
            todos_registos.setdefault(provincia, []).extend(registos)
            print(f'  {provincia}: +{len(registos)} registos lidos de "{nome_ficheiro}"')

    print()
    for provincia, registos in todos_registos.items():
        caminho_csv = gravar_csv(provincia, registos)
        print(f'OK {provincia}: {len(registos)} registos escritos em {os.path.relpath(caminho_csv, BASE_DIR)}')


if __name__ == '__main__':
    caminho = sys.argv[1] if len(sys.argv) > 1 else PASTA_ORIGEM_PADRAO
    if not os.path.exists(caminho):
        print(f'Caminho não encontrado: {caminho}')
        print(f'Coloca as matrizes .xlsx em: {PASTA_ORIGEM_PADRAO}')
        sys.exit(1)
    converter(caminho)
