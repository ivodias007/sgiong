#!/usr/bin/env python3
"""
converter_xls_para_xlsx.py — Converte ficheiros .xls (Excel 97-2003) para
.xlsx, célula a célula, para poderem ser lidos por openpyxl (usado por
converter_excel_para_csv.py e importar_novas_matrizes.py).

Necessário para: MATRIZ ONGs SOFALA.xls, Provincia de Maputo Matriz.xls,
Zambezia- Matriz.xls — que vieram no formato antigo.

Uso:
  pip install xlrd openpyxl --break-system-packages
  python converter_xls_para_xlsx.py origem_excel/MATRIZ ONGs SOFALA.xls
  # ou, para converter todos os .xls de uma pasta:
  python converter_xls_para_xlsx.py origem_excel/
"""

import sys
import os
import glob
import xlrd
import openpyxl


def converter(caminho_xls):
    livro_antigo = xlrd.open_workbook(caminho_xls, formatting_info=False)
    livro_novo = openpyxl.Workbook()
    livro_novo.remove(livro_novo.active)

    for nome_folha in livro_antigo.sheet_names():
        folha_antiga = livro_antigo.sheet_by_name(nome_folha)
        folha_nova = livro_novo.create_sheet(title=nome_folha[:31])  # limite do Excel
        for r in range(folha_antiga.nrows):
            for c in range(folha_antiga.ncols):
                valor = folha_antiga.cell_value(r, c)
                if valor != '':
                    folha_nova.cell(row=r + 1, column=c + 1, value=valor)

    caminho_novo = os.path.splitext(caminho_xls)[0] + '.xlsx'
    livro_novo.save(caminho_novo)
    return caminho_novo


def main():
    if len(sys.argv) < 2:
        print('Uso: python converter_xls_para_xlsx.py <ficheiro.xls | pasta>')
        sys.exit(1)

    alvo = sys.argv[1]
    ficheiros = glob.glob(os.path.join(alvo, '*.xls')) if os.path.isdir(alvo) else [alvo]
    ficheiros = [f for f in ficheiros if not f.lower().endswith('.xlsx')]

    for f in ficheiros:
        novo = converter(f)
        print(f'OK: {f} -> {novo}')


if __name__ == '__main__':
    main()
