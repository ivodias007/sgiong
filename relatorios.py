"""
relatorios.py — Geração de relatórios PDF, Excel e CSV do SGIONGs.

Todos os relatórios respeitam os mesmos filtros da listagem/dashboard
(província, área, distrito, natureza, origem, período) e só incluem ONGs
aprovadas — nunca pendentes, rejeitadas ou rascunhos.
"""

import io
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                 PageBreak)
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import HorizontalBarChart

VERDE = colors.HexColor('#0B4228')
VERDE_ESCURO = colors.HexColor('#082E1C')
TERRACOTA = colors.HexColor('#B5451B')
DOURADO = colors.HexColor('#B08A3E')
PAPEL = colors.HexColor('#FBF9F5')

MESES_PT = ['', 'Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho', 'Julho',
            'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']

COLUNAS_EXPORTACAO = [
    ('nome', 'Nome'), ('sigla', 'Sigla'), ('natureza', 'Natureza'), ('origem', 'Origem'),
    ('provincia', 'Província'), ('area_intervencao', 'Área de Intervenção'),
    ('distritos_texto', 'Distritos'), ('ano_entrada_pais', 'Ano de Entrada'),
    ('pessoa_contacto', 'Ponto Focal'), ('telefone', 'Telefone'), ('email', 'E-mail'),
    ('fonte_financiamento', 'Fonte de Financiamento'), ('estado', 'Estado'),
    ('criado_em', 'Registada em'),
]


def _buscar_ongs(conn, filtros):
    """Devolve as ONGs aprovadas que respeitam os filtros — mesma lógica de
    âmbito que a listagem principal, para os relatórios nunca mostrarem mais
    do que a pessoa já podia ver no site."""
    condicoes = ["o.aprovacao = 'Aprovado'"]
    params = []
    joins = ''

    if filtros.get('provincia'):
        joins += ' JOIN ong_provincias opf ON opf.ong_id = o.id '
        condicoes.append('opf.provincia = ?')
        params.append(filtros['provincia'])
    if filtros.get('area_id'):
        joins += ' JOIN ong_areas oaf ON oaf.ong_id = o.id '
        condicoes.append('oaf.area_id = ?')
        params.append(filtros['area_id'])
    if filtros.get('distrito_id'):
        joins += ' JOIN ong_distritos odf ON odf.ong_id = o.id '
        condicoes.append('odf.distrito_id = ?')
        params.append(filtros['distrito_id'])
    if filtros.get('natureza'):
        condicoes.append('o.id IN (SELECT ong_id FROM ong_naturezas WHERE natureza = ?)')
        params.append(filtros['natureza'])
    if filtros.get('origem'):
        condicoes.append('o.origem = ?')
        params.append(filtros['origem'])
    if filtros.get('desde'):
        condicoes.append('date(o.criado_em) >= date(?)')
        params.append(filtros['desde'])
    if filtros.get('ate'):
        condicoes.append('date(o.criado_em) <= date(?)')
        params.append(filtros['ate'])

    query = f'SELECT DISTINCT o.* FROM ongs o {joins} WHERE {" AND ".join(condicoes)} ORDER BY o.nome'
    return conn.execute(query, params).fetchall()


def _resumo_indicadores(conn, filtros, ongs):
    """Indicadores calculados SEMPRE sobre o mesmo conjunto de ONGs que aparece na
    listagem do relatório (ou seja, respeitam os filtros) — assim o resumo, os
    gráficos e a listagem nunca se contradizem."""
    ids = [o['id'] for o in ongs]
    total = len(ids)
    por_provincia, por_area = {}, {}
    if ids:
        marcas = ','.join('?' * len(ids))
        for r in conn.execute(f'''
            SELECT provincia, COUNT(DISTINCT ong_id) c FROM ong_provincias
            WHERE ong_id IN ({marcas}) GROUP BY provincia ORDER BY c DESC, provincia
        ''', ids):
            por_provincia[r['provincia']] = r['c']
        for r in conn.execute(f'''
            SELECT a.nome, COUNT(DISTINCT oa.ong_id) c FROM ong_areas oa
            JOIN areas a ON a.id = oa.area_id
            WHERE oa.ong_id IN ({marcas}) GROUP BY a.nome ORDER BY c DESC, a.nome
        ''', ids):
            por_area[r['nome']] = r['c']

    nacionais = sum(1 for o in ongs if (o['origem'] or '') == 'Nacional')
    estrangeiras = sum(1 for o in ongs if (o['origem'] or '') == 'Estrangeira')

    return {
        'total': total,
        'nacionais': nacionais,
        'estrangeiras': estrangeiras,
        'provincias_cobertas': len(por_provincia),
        'por_provincia': por_provincia,
        'por_area': por_area,
    }


# ==========================================================================
# CSV
# ==========================================================================
def gerar_csv(conn, filtros):
    ongs = _buscar_ongs(conn, filtros)
    saida = io.StringIO()
    escritor = csv.writer(saida)
    escritor.writerow([rotulo for _, rotulo in COLUNAS_EXPORTACAO])
    for o in ongs:
        escritor.writerow([o[campo] if o[campo] is not None else '' for campo, _ in COLUNAS_EXPORTACAO])
    return saida.getvalue().encode('utf-8-sig')  # BOM para abrir bem acentos no Excel


# ==========================================================================
# EXCEL
# ==========================================================================
def gerar_excel(conn, filtros):
    ongs = _buscar_ongs(conn, filtros)
    resumo = _resumo_indicadores(conn, filtros, ongs)

    wb = Workbook()

    # --- Folha 1: Resumo ---
    ws1 = wb.active
    ws1.title = 'Resumo'
    fundo_verde = PatternFill(start_color='0B4228', end_color='0B4228', fill_type='solid')
    fonte_branca_negrito = Font(color='FFFFFF', bold=True, size=12)

    ws1['A1'] = 'Relatório Nacional das ONGs'
    ws1['A1'].font = Font(bold=True, size=16, color='0B4228')
    ws1['A2'] = f'Gerado em {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws1['A2'].font = Font(italic=True, size=9, color='52514C')

    linha = 4
    for rotulo, valor in [
        ('Total de ONGs (filtro aplicado)', resumo['total']),
        ('Nacionais', resumo['nacionais']),
        ('Estrangeiras', resumo['estrangeiras']),
        ('Províncias cobertas', resumo['provincias_cobertas']),
    ]:
        ws1.cell(row=linha, column=1, value=rotulo).font = Font(bold=True)
        ws1.cell(row=linha, column=2, value=valor)
        linha += 1

    linha += 1
    ws1.cell(row=linha, column=1, value='Distribuição por província').font = fonte_branca_negrito
    ws1.cell(row=linha, column=1).fill = fundo_verde
    ws1.cell(row=linha, column=2).fill = fundo_verde
    linha += 1
    for provincia, total in resumo['por_provincia'].items():
        ws1.cell(row=linha, column=1, value=provincia)
        ws1.cell(row=linha, column=2, value=total)
        linha += 1

    linha += 1
    ws1.cell(row=linha, column=1, value='Distribuição por área temática').font = fonte_branca_negrito
    ws1.cell(row=linha, column=1).fill = fundo_verde
    ws1.cell(row=linha, column=2).fill = fundo_verde
    linha += 1
    for area, total in resumo['por_area'].items():
        ws1.cell(row=linha, column=1, value=area)
        ws1.cell(row=linha, column=2, value=total)
        linha += 1

    ws1.column_dimensions['A'].width = 38
    ws1.column_dimensions['B'].width = 14

    # --- Folha 2: Listagem completa ---
    ws2 = wb.create_sheet('ONGs')
    for col, (_, rotulo) in enumerate(COLUNAS_EXPORTACAO, start=1):
        celula = ws2.cell(row=1, column=col, value=rotulo)
        celula.font = fonte_branca_negrito
        celula.fill = fundo_verde
        celula.alignment = Alignment(horizontal='left')

    for linha_idx, o in enumerate(ongs, start=2):
        for col, (campo, _) in enumerate(COLUNAS_EXPORTACAO, start=1):
            ws2.cell(row=linha_idx, column=col, value=o[campo])

    for col in range(1, len(COLUNAS_EXPORTACAO) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 22
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = f'A1:{get_column_letter(len(COLUNAS_EXPORTACAO))}{len(ongs) + 1}'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ==========================================================================
# PDF
# ==========================================================================
def _grafico_barras_horizontais(dados, largura=420, altura=180, cor=TERRACOTA):
    """Pequeno gráfico de barras horizontais (top 10), usando reportlab.graphics
    — sem depender de matplotlib, mais leve para o servidor."""
    itens = list(dados.items())[:10]
    desenho = Drawing(largura, altura)
    grafico = HorizontalBarChart()
    grafico.x = 120
    grafico.y = 10
    grafico.height = altura - 20
    grafico.width = largura - 140
    grafico.data = [[v for _, v in reversed(itens)]]
    grafico.categoryAxis.categoryNames = [k for k, _ in reversed(itens)]
    grafico.categoryAxis.labels.fontSize = 7
    grafico.valueAxis.valueMin = 0
    grafico.bars[0].fillColor = cor
    grafico.barLabels.fontSize = 7
    grafico.barLabelFormat = '%d'
    grafico.bars.strokeWidth = 0
    desenho.add(grafico)
    return desenho


def gerar_pdf(conn, filtros, titulo_contexto=None):
    ongs = _buscar_ongs(conn, filtros)
    resumo = _resumo_indicadores(conn, filtros, ongs)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                             topMargin=2 * cm, bottomMargin=2 * cm,
                             leftMargin=1.8 * cm, rightMargin=1.8 * cm)

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle('TituloCapa', parent=estilos['Title'], fontSize=26,
                                   textColor=VERDE_ESCURO, spaceAfter=6, alignment=1)
    estilo_subtitulo = ParagraphStyle('Subtitulo', parent=estilos['Normal'], fontSize=13,
                                      textColor=TERRACOTA, alignment=1, spaceAfter=4)
    estilo_institucional = ParagraphStyle('Institucional', parent=estilos['Normal'], fontSize=10,
                                          textColor=colors.HexColor('#52514C'), alignment=1)
    estilo_secao = ParagraphStyle('Seccao', parent=estilos['Heading2'], textColor=VERDE_ESCURO,
                                  spaceBefore=14, spaceAfter=8)
    estilo_corpo = ParagraphStyle('Corpo', parent=estilos['Normal'], fontSize=9)

    elementos = []

    # ---- Capa ----
    elementos.append(Spacer(1, 4 * cm))
    elementos.append(Paragraph('República de Moçambique', estilo_institucional))
    elementos.append(Paragraph('Ministério da Planificação e Desenvolvimento', estilo_institucional))
    elementos.append(Spacer(1, 1.5 * cm))
    mes_nome = MESES_PT[datetime.now().month]
    elementos.append(Paragraph('Relatório Nacional das ONGs', estilo_titulo))
    elementos.append(Paragraph(f'{mes_nome} de {datetime.now().year}', estilo_subtitulo))
    if titulo_contexto:
        elementos.append(Paragraph(titulo_contexto, estilo_subtitulo))
    elementos.append(Spacer(1, 3 * cm))
    elementos.append(Paragraph(f'Gerado automaticamente pelo SGIONGs em {datetime.now().strftime("%d/%m/%Y às %H:%M")}',
                                estilo_institucional))
    elementos.append(PageBreak())

    # ---- Indicadores-chave ----
    elementos.append(Paragraph('Indicadores-chave', estilo_secao))
    dados_indicadores = [
        ['Total de ONGs', 'Nacionais', 'Estrangeiras', 'Províncias cobertas'],
        [str(resumo['total']), str(resumo['nacionais']), str(resumo['estrangeiras']), str(resumo['provincias_cobertas'])],
    ]
    tabela_ind = Table(dados_indicadores, colWidths=[4 * cm] * 4)
    tabela_ind.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), VERDE_ESCURO),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 1), (-1, 1), 16),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('TEXTCOLOR', (0, 1), (-1, 1), VERDE_ESCURO),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E4E0D8')),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    elementos.append(tabela_ind)

    # ---- Gráfico por província ----
    if resumo['por_provincia']:
        elementos.append(Paragraph('Distribuição por província', estilo_secao))
        elementos.append(_grafico_barras_horizontais(resumo['por_provincia'], cor=TERRACOTA))

    # ---- Gráfico por área temática ----
    if resumo['por_area']:
        elementos.append(Paragraph('Distribuição por área temática', estilo_secao))
        elementos.append(_grafico_barras_horizontais(resumo['por_area'], cor=DOURADO))

    elementos.append(PageBreak())

    # ---- Tabela completa (até 200 linhas, para o PDF não ficar gigante) ----
    elementos.append(Paragraph(f'Listagem de organizações ({len(ongs)} no total'
                                f'{", mostrando as primeiras 200" if len(ongs) > 200 else ""})',
                                estilo_secao))
    cabecalho = ['Nome', 'Natureza', 'Origem', 'Província', 'Área de Intervenção']
    linhas_tabela = [cabecalho]
    for o in ongs[:200]:
        linhas_tabela.append([
            Paragraph(o['nome'] or '', estilo_corpo),
            o['natureza'] or '—',
            o['origem'] or '—',
            o['provincia'] or '—',
            Paragraph((o['area_intervencao'] or '—')[:80], estilo_corpo),
        ])
    tabela_ongs = Table(linhas_tabela, colWidths=[5 * cm, 2 * cm, 2.2 * cm, 2.8 * cm, 5 * cm], repeatRows=1)
    tabela_ongs.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), VERDE),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#E4E0D8')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, PAPEL]),
    ]))
    elementos.append(tabela_ongs)

    def rodape(canvas, doc_):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#52514C'))
        canvas.drawString(1.8 * cm, 1.3 * cm, 'SGIONGs — Ministério da Planificação e Desenvolvimento')
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.3 * cm, f'Página {doc_.page}')
        canvas.restoreState()

    doc.build(elementos, onFirstPage=rodape, onLaterPages=rodape)
    buffer.seek(0)
    return buffer.getvalue()
