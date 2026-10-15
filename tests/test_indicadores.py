"""Fiabilidade dos indicadores e dos relatórios: mesmos números em todo o lado."""
import csv
import io

import openpyxl
import pytest
from flask import template_rendered
from pypdf import PdfReader

import app as m
from conftest import cliente_com, texto
from tagging import PROVINCIAS


@pytest.fixture
def ctx():
    guardado = {}
    def rec(sender, template, context, **kw):
        guardado[template.name] = context
    template_rendered.connect(rec, m.app)
    yield guardado
    template_rendered.disconnect(rec, m.app)


def sql_total(db, **f):
    q = "SELECT COUNT(DISTINCT o.id) FROM ongs o WHERE o.aprovacao='Aprovado' AND o.rascunho=0"
    p = []
    if f.get('provincia'):
        q += ' AND o.id IN (SELECT ong_id FROM ong_provincias WHERE provincia=?)'; p.append(f['provincia'])
    if f.get('distrito_id'):
        q += ' AND o.id IN (SELECT ong_id FROM ong_distritos WHERE distrito_id=?)'; p.append(f['distrito_id'])
    if f.get('area_id'):
        q += ' AND o.id IN (SELECT ong_id FROM ong_areas WHERE area_id=?)'; p.append(f['area_id'])
    if f.get('natureza'):
        q += ' AND o.id IN (SELECT ong_id FROM ong_naturezas WHERE natureza=?)'; p.append(f['natureza'])
    if f.get('origem'):
        q += ' AND o.origem=?'; p.append(f['origem'])
    return db.execute(q, p).fetchone()[0]


def linhas_csv(c, qs):
    r = c.get('/relatorios/csv' + qs)
    return list(csv.reader(io.StringIO(texto(r).lstrip('﻿'))))[1:]


def linhas_excel(c, qs):
    wb = openpyxl.load_workbook(io.BytesIO(c.get('/relatorios/excel' + qs).data))
    ws = wb.worksheets[-1] if len(wb.worksheets) > 1 else wb.worksheets[0]
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]], wb


def texto_pdf(c, qs):
    r = c.get('/relatorios/pdf' + qs)
    assert r.status_code == 200 and r.data[:4] == b'%PDF'
    return '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(r.data)).pages)


@pytest.mark.parametrize('prov', PROVINCIAS)
def test_total_por_provincia_igual_em_todo_o_sistema(db, ctx, prov):
    c = cliente_com('super_admin')
    esperado = sql_total(db, provincia=prov)
    c.get(f'/dashboard?provincia={prov}')
    assert ctx['dashboard.html']['total_ongs'] == esperado, 'painel'
    c.get(f'/ongs?provincia={prov}')
    assert ctx['index.html']['total'] == esperado, 'listagem'
    assert len(linhas_csv(c, f'?provincia={prov}')) == esperado, 'CSV'
    mapa = {d['nome']: d['total'] for d in ctx['dashboard.html']['mapa_dados']} if False else None
    c.get('/dashboard')
    assert {d['nome']: d['total'] for d in ctx['dashboard.html']['mapa_dados']}.get(prov, 0) == esperado, 'mapa'


def test_total_nacional_igual_em_todo_o_sistema(db, ctx):
    c = cliente_com('super_admin')
    esperado = sql_total(db)
    c.get('/dashboard'); assert ctx['dashboard.html']['total_ongs'] == esperado
    c.get('/ongs'); assert ctx['index.html']['total'] == esperado
    assert len(linhas_csv(c, '')) == esperado
    assert len(linhas_excel(c, '')[0]) == esperado
    assert f'{esperado} no total' in texto_pdf(c, '')


@pytest.mark.parametrize('qs,filtros', [
    ('?provincia=Sofala&natureza=ONG', dict(provincia='Sofala', natureza='ONG')),
    ('?natureza=OSC', dict(natureza='OSC')),
    ('?origem=Estrangeira', dict(origem='Estrangeira')),
    ('?provincia=Tete&origem=Nacional&natureza=ONG', dict(provincia='Tete', origem='Nacional', natureza='ONG')),
])
def test_filtros_combinados_csv_excel_pdf_iguais(db, qs, filtros):
    c = cliente_com('super_admin')
    esperado = sql_total(db, **filtros)
    assert esperado > 0
    assert len(linhas_csv(c, qs)) == esperado
    assert len(linhas_excel(c, qs)[0]) == esperado
    assert f'{esperado} no total' in texto_pdf(c, qs)


def test_filtro_area_e_distrito(db):
    c = cliente_com('super_admin')
    area = db.execute('SELECT area_id, COUNT(*) n FROM ong_areas GROUP BY area_id ORDER BY n DESC').fetchone()[0]
    dist = db.execute('SELECT distrito_id, COUNT(*) n FROM ong_distritos GROUP BY distrito_id ORDER BY n DESC').fetchone()[0]
    for f in (dict(area_id=area), dict(distrito_id=dist), dict(area_id=area, provincia='Sofala')):
        qs = '?' + '&'.join(f'{k}={v}' for k, v in f.items())
        assert len(linhas_csv(c, qs)) == sql_total(db, **f)


def test_resumo_dos_relatorios_respeita_o_filtro(db):
    """Os gráficos/resumo por província e área usam só as ONGs filtradas (antes mostravam o país inteiro)."""
    c = cliente_com('super_admin')
    _, wb = linhas_excel(c, '?provincia=Tete&natureza=ONG')
    resumo = wb.worksheets[0]
    celulas = [[x for x in r] for r in resumo.iter_rows(values_only=True)]
    total = [r[1] for r in celulas if r[0] and 'Total de ONGs' in str(r[0])][0]
    assert total == sql_total(db, provincia='Tete', natureza='ONG')
    por_prov = {r[0]: r[1] for r in celulas if r[0] in PROVINCIAS}
    assert por_prov.get('Tete') == total      # todas as ONGs filtradas actuam em Tete
    assert all(v <= total for v in por_prov.values())


def test_origem_nacional_mais_estrangeira_nao_excede_total(db, ctx):
    c = cliente_com('super_admin')
    c.get('/dashboard')
    d = ctx['dashboard.html']
    assert sum(d['origem_valores']) == d['total_ongs']          # inclui o balde "Não especificado"
    assert sum(d['estado_valores']) == d['total_ongs']


def test_totais_do_painel_acompanham_aprovacao(db, ctx):
    """Uma ONG que passa a pendente/rejeitada sai de todos os totais."""
    c = cliente_com('super_admin')
    antes = sql_total(db)
    oid = db.execute("SELECT id FROM ongs WHERE aprovacao='Aprovado' LIMIT 1").fetchone()[0]
    db.execute("UPDATE ongs SET aprovacao='Rejeitado' WHERE id=?", (oid,)); db.commit()
    c.get('/dashboard'); assert ctx['dashboard.html']['total_ongs'] == antes - 1
    from conftest import Cliente
    Cliente().get('/ongs'); assert ctx['index.html']['total'] == antes - 1     # vista pública: só aprovadas
    assert len(linhas_csv(c, '')) == antes - 1


def test_gestor_provincial_ve_numeros_da_sua_provincia(db, ctx):
    c = cliente_com('gestor_provincial', 'Zambézia')
    c.get('/dashboard?provincia=Sofala')
    assert ctx['dashboard.html']['total_ongs'] == sql_total(db, provincia='Zambézia')


def test_listagem_por_distrito_bate_certo(db, ctx):
    c = cliente_com('super_admin')
    d = db.execute('''SELECT d.id, d.provincia FROM distritos d JOIN ong_distritos od ON od.distrito_id=d.id
                      GROUP BY d.id ORDER BY COUNT(*) DESC LIMIT 1''').fetchone()
    c.get(f'/ongs?provincia={d["provincia"]}&distrito_id={d["id"]}')
    assert ctx['index.html']['total'] == sql_total(db, provincia=d['provincia'], distrito_id=d['id'])
